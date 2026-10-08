# -*- coding: utf-8 -*-
"""يولّد نسختي Word وExcel من نفس محتوى مرجع IQ."""
import os
import sys
import re
import math

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
sys.path.insert(0, HERE)

from svg2png import render_all  # noqa: E402
import content_a, content_b, content_c, content_d, content_e  # noqa: E402

CHAPTERS = content_a.CH + content_b.CH + content_c.CH + content_d.CH + content_e.CH

SOURCES = [
    ('بسهولة شرح أسئلة IQ (نسبة الذكاء) - الجزء الأول', 'V_zkkehBvbQ', 'مصفوفات وأشكال'),
    ('بسهولة شرح أسئلة IQ - الشكل المكمل - الجزء الثاني', 'wtUI5X2ppjs', 'الشكل المكمل'),
    ('بسهولة شرح أسئلة IQ - الجزء الثاني', 'bjytZM1dJOw', 'أشكال متنوعة'),
    ('بسهولة شرح أسئلة IQ (المكعبات) - الجزء الرابع', '4Qk4AJJp7Yo', 'المكعبات'),
    ('بسهولة شرح أسئلة IQ (المصفوفات) - الجزء الخامس', 'hfRxRoU8OFY', 'المصفوفات'),
    ('بسهولة شرح أسئلة IQ (المرايا وثني الورق) - الجزء السادس', 'MDt1D3NzB24', 'المرايا وثني الورق'),
    ('بسهولة شرح أسئلة IQ (تعبيرات الوجه) - الجزء السابع', 'Ouyv9RTYXHQ', 'تعبيرات الوجه'),
    ('شرح أسئلة IQ (الأسئلة المقالية) - الجزء الثامن', '5t8OAwQMEFk', 'الأسئلة المقالية'),
    ('حل IQ اختبار تدوير الأشكال - اختبار بيردو للتصور المكاني', 'D9kUXTcia4A', 'التصور المكاني'),
    ('أحدث أسئلة IQ - المتسلسلة - المتتاليات - المتواليات', '-x_B02vr9QY', 'المتتاليات'),
    ('شرح أحدث أسئلة IQ - حل أشكال X, Y, Z بسهولة', '2Ihx_bTZzTU', 'تركيب الأشكال'),
    ('شرح متتاليات بالكسور العشرية', '3uY8YdcW0ic', 'الكسور العشرية'),
    ('متتاليات الحروف - أحدث أسئلة IQ', 'cAIyFyMa7VY', 'متتاليات الحروف'),
    ('IQ - شرح أسئلة تركيب الأشكال بمعلومية X , Y , Z', '5jx2zlqbyC8', 'تركيب الأشكال'),
    ('شرح أحدث أشكال IQ - وردت بالفعل في الاختبارات', 'YoEJZqi61CM', 'أشكال حديثة'),
    ('حل أشكال IQ الحديثة التي وردت بالفعل خلال الفترة السابقة', '2oSouwEyWZE', 'أشكال حديثة'),
    ('شرح أحدث أسئلة IQ بسهولة - تخصص لغة عربية 2024', 'H7oXyOwewpM', 'مراجعة مطولة'),
    ('أحدث أسئلة IQ (طي الورق) وردت بالفعل في الاختبارات', 'i6nL9sjz-Dw', 'طي الورق'),
    ('شرح أحدث أسئلة IQ - المتتابعات والمتسلسلة', 'RIMlB-lU3Zs', 'المتتابعات'),
    ('أحدث أشكال نسبة الذكاء IQ وردت بالفعل - 2025', 'fE9INmOeuGs', 'أشكال 2025'),
    ('شرح أحدث أسئلة IQ - ورقة سمكة يد (ألغاز الأشكال)', 'pZWlqSyV5qA', 'ألغاز الأشكال'),
    ('شرح أسئلة ونماذج الميزان من محور IQ', 'x7G9D552NRc', 'الميزان'),
    ('أحدث أسئلة IQ - العلاقة بين الأشياء (الشوكة والملعقة...)', 'lpjD8DxhEKk', 'العلاقة بين الأشياء'),
    ('شرح أسئلة IQ - الأمثال الشعبية ولماذا نستخدم؟', 'YtyenRbHXkA', 'الأمثال الشعبية'),
    ('أسهل طريقة لحل طي المكعبات - المكعبات المفرودة', 'VhTX7E_XJNA', 'طي المكعبات'),
    ('أسهل طريقة لحل IQ - استخرج الشكل الخاطئ للمكعبات', 'tR_-va2vlmA', 'استراتيجية عقارب الساعة'),
    ('بأسهل طرق الحل - دليلك الشامل لحل طي المكعبات', 'nq1hV7tput0', 'طي المكعبات'),
    ('شرح أحدث أسئلة المتتاليات (المتتابعات) من محور IQ', 'PUF2tQlIJFk', 'المتتاليات'),
    ('مراجعة ما قبل الامتحان - العلاقة بين الأشياء - الأمثال', 'NFJ627OHhfQ', 'التفكير المجرد'),
    ('أحدث أسئلة IQ العلاقات اللفظية - وردت بالفعل', 'dz-sE6m2qfg', 'العلاقات اللفظية'),
    ('بسهولة حل أسئلة IQ - المكعبات والعلاقة بين الأشياء', 'CR5hlBRX6Go', 'مراجعة'),
    ('أحدث أسئلة في IQ - طريقة حل أسئلة نسبة الذكاء 2026', 'XbokmztXcuw', 'مراجعة 2026'),
    ('أحدث أسئلة في IQ - مراجعة ما قبل الامتحان', 'SoeGN11cUAQ', 'مراجعة'),
    ('حل أحدث مسألة من أشكال IQ بطريقة سهلة وبسيطة', 'oIpba1Dz4OY', 'مصفوفة أشكال'),
    ('أحدث أسئلة في IQ - لجميع الوظائف 2025', 'kK_LDGJVMwU', 'مراجعة 2025'),
    ('أحدث أشكال IQ وردت بالفعل - مسابقة الأزهر الشريف', 'zapEEiuKolw', 'الأزهر'),
    ('شرح ومراجعة IQ - الأشكال المماثلة - العلاقات - تدوير المكعب', 'h_d5Q6PTVPk', 'مراجعة'),
    ('التناظر اللفظي في IQ - العلاقة بين الكلمات', 'rdee9dRkE7g', 'التناظر اللفظي'),
    ('أحدث أسئلة وردت بالفعل في IQ - المماثل + المرايا + المتتاليات', '3zT0e7wN5Jw', 'مراجعة'),
    ('حل أحدث أسئلة IQ والتفكير المنطقي - 2026', 'Nm4w4T-SQ4U', 'تفكير منطقي'),
    ('أحدث أشكال IQ - مسابقة الطب البيطري', 'BVUNf8AFYHU', 'الطب البيطري'),
    ('أحدث أشكال IQ - تركيب الأشكال XYZ - المتماثلة 2026', '68GlrljDVXk', 'مراجعة 2026'),
    ('شرح أحدث أشكال IQ - تخصص الدراسات بمسابقة الحصة', 'rXsi_P42sBA', 'تخصص الدراسات'),
    ('شكل IQ من (منسا) حيّر الجميع - أحدث أشكال IQ', '0bz1oMT6zq0', 'شكل منسا'),
    ('حل أكثر من 10 أشكال IQ من أصعب الأشكال', 'OmU6XDh_peI', 'أشكال صعبة'),
    ('أحدث أنماط IQ - تخصص اللغة الإنجليزية 2026', 'qm-HIrolrQI', 'أنماط 2026'),
    ('أحدث 10 أشكال IQ - الأزهر والأحساء 2026', 'Cz-tTLKItEc', 'مراجعة 2026'),
    ('أهم 20 شكل IQ وردت بالفعل في اختبارات التوظيف', 'LYWSbDP100o', '20 شكلًا'),
    ('مسابقة الخبراء - أهم أنماط IQ وردت بالفعل', 'yOJd0tAVxIk', 'مسابقة الخبراء'),
    ('بأسهل طرق الحل - المتطابقة والمماثلة وتدوير الأشكال', 'pmDaEsMNvkI', 'المتطابقة والمماثلة'),
]

PLAN = [
    ("اليوم 1", "المتتاليات + المصفوفات الرقمية", "احفظ أنماط الفروق والمربعات والمكعبات، وحُل 30 مسألة."),
    ("اليوم 2", "متتاليات الحروف + الشكل المكمل", "اكتب ترتيب الحروف على ورقة، وحل السلسلتين بخطوتين."),
    ("اليوم 3", "المكعبات", "الاستبعاد - المثلثات - عقارب الساعة (بمكعب ورقي حقيقي)."),
    ("اليوم 4", "التصور المكاني + التدوير (بيردو)", "استخدم علبة ثقاب أو مكعبًا ودوّره بيدك."),
    ("اليوم 5", "الأشكال المتطابقة + تركيب X/Y/Z", "ركز على الفرق بين التدوير والانعكاس."),
    ("اليوم 6", "المرايا وثني الورق + تعبيرات الوجه", "جرب ثني الورق بيدك، واحفظ علامات المشاعر السبعة."),
    ("اليوم 7", "مراجعة شاملة + اختبار زمني", "30 سؤالًا في 25 دقيقة."),
]

TIPS = [
    "اقرأ رأس السؤال: هل المطلوب الصحيح أم «الخاطئ» أم «المختلف»؟",
    "لا تتوقف أكثر من 45 ثانية على السؤال — علّمه وارجع إليه.",
    "استخدم الاستبعاد قبل التخيّل: حذف اختيارين يرفع إصابتك من 25% إلى 50%.",
    "في الأشكال: سمِّ الأشكال بصوت منخفض (مثلث، دائرة، مربع).",
    "اكتب الفروق بين الأرقام فوق السلسلة مباشرة.",
    "راجع في آخر 5 دقائق: تأكد أنك لم تختر «الصحيح» في سؤال يطلب «الخاطئ».",
    "التدريب أهم من الحفظ: نفس الأسئلة تتكرر بأشكال قريبة جدًا.",
]

COLOR = {c["key"]: c["color"] for c in CHAPTERS}
TOTAL_Q = sum(len(c["qs"]) for c in CHAPTERS)

TAG = re.compile(r"<[^>]+>")


def plain(s):
    return TAG.sub("", s).replace("&nbsp;", " ").strip()


def figures(scale=2.0):
    outdir = os.path.join(ROOT, "assets", "figs")
    return render_all(CHAPTERS, outdir, scale=scale)


# ============================================================== WORD
def build_docx(figs, out_path):
    from docx import Document
    from docx.shared import Pt, Cm, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement

    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.left_margin = sec.right_margin = Cm(1.6)
    sec.top_margin = sec.bottom_margin = Cm(1.6)

    normal = doc.styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(12)
    normal._element.rPr.rFonts.set(qn("w:cs"), "Arial")

    def rtl(p, align="right"):
        pPr = p._p.get_or_add_pPr()
        bidi = OxmlElement("w:bidi")
        bidi.set(qn("w:val"), "1")
        pPr.append(bidi)
        p.alignment = {"right": WD_ALIGN_PARAGRAPH.RIGHT,
                       "center": WD_ALIGN_PARAGRAPH.CENTER,
                       "left": WD_ALIGN_PARAGRAPH.LEFT}[align]

    def shade(el, hexcolor):
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear")
        shd.set(qn("w:color"), "auto")
        shd.set(qn("w:fill"), hexcolor)
        el.append(shd)

    def para(text="", size=12, bold=False, color=None, align="right", space_after=4):
        p = doc.add_paragraph()
        rtl(p, align)
        p.paragraph_format.space_after = Pt(space_after)
        if text:
            r = p.add_run(text)
            r.font.size = Pt(size)
            r.font.bold = bold
            r.font.name = "Arial"
            r._element.rPr.rFonts.set(qn("w:cs"), "Arial")
            if color:
                r.font.color.rgb = RGBColor.from_string(color.lstrip("#"))
        return p

    def rich(segments, size=12, align="right"):
        """segments: list of (text, bold, color)"""
        p = doc.add_paragraph()
        rtl(p, align)
        p.paragraph_format.space_after = Pt(3)
        for text, bold, color in segments:
            r = p.add_run(text)
            r.font.size = Pt(size)
            r.font.bold = bold
            r.font.name = "Arial"
            r._element.rPr.rFonts.set(qn("w:cs"), "Arial")
            if color:
                r.font.color.rgb = RGBColor.from_string(color.lstrip("#"))
        return p

    def cell_shade(cell, hexcolor):
        shade(cell._tc.get_or_add_tcPr(), hexcolor)

    def table_rtl(table):
        tblPr = table._tbl.tblPr
        bidi = OxmlElement("w:bidi")
        bidi.set(qn("w:val"), "1")
        tblPr.append(bidi)

    # ---------------- cover
    p = para("المرجع الشامل لأسئلة القدرات الذهنية IQ", size=26, bold=True,
             color="#2a6fdb", align="center", space_after=6)
    shade(p._p.get_or_add_pPr(), "EAF1FF")
    para("جمع وتلخيص لكل أسئلة IQ (الصور والأشكال والأعداد والحروف) مع طريقة الحل خطوة بخطوة",
         size=13, align="center", color="#5b6b8a", space_after=4)
    para("مستخلص من قناة «هنتعلم أون لاين مع إيمان السيد» — مع رسم توضيحي لكل سؤال بجانب الشرح",
         size=12, align="center", color="#7b4bd8", bold=True, space_after=10)
    para(f"{len(CHAPTERS)} بابًا  •  {TOTAL_Q} سؤالًا مشروحًا  •  {len(SOURCES)} فيديو مصدر  •  88 رسمًا توضيحيًا",
         size=12, bold=True, align="center", color="#2a9d8f", space_after=10)
    pn = para("تنبيه: الرسوم إعادة رسم توضيحية مبنية على شرح الفيديوهات (لا يمكن أخذ لقطات من الفيديو)، "
              "وهي مطابقة في الفكرة والخصائص للتدريب. المحتوى استرشادي.",
              size=10.5, align="center", color="#b26a12", space_after=2)
    shade(pn._p.get_or_add_pPr(), "FFF8E6")

    # ---------------- toc
    para("فهرس الأبواب", size=18, bold=True, color="#1f2a44", space_after=6)
    for i, c in enumerate(CHAPTERS, 1):
        rich([(f"{i}. ", True, c["color"]), (f'{c["icon"]}  {c["title"]}', True, c["color"]),
              (f'  —  {len(c["qs"])} سؤالًا', False, "#5b6b8a")], size=12)

    doc.add_page_break()

    # ---------------- chapters
    for ci, c in enumerate(CHAPTERS, 1):
        hp = para(f'{c["icon"]}  الباب {ci}: {c["title"]}', size=19, bold=True,
                  color="FFFFFF", space_after=6)
        shade(hp._p.get_or_add_pPr(), c["color"].lstrip("#"))
        para(plain(c["intro"]), size=12, space_after=6)

        rp = para("قواعد الحل السريع", size=14, bold=True, color=c["color"], space_after=3)
        shade(rp._p.get_or_add_pPr(), "F1F5FF")
        for r in c["rules"]:
            para("• " + plain(r), size=11.5, space_after=2)

        for qi, q in enumerate(c["qs"], 1):
            qp = para(f'س{qi} — {q["title"]}', size=14, bold=True, color="FFFFFF", space_after=4)
            shade(qp._p.get_or_add_pPr(), c["color"].lstrip("#"))

            table = doc.add_table(rows=1, cols=2)
            table.alignment = WD_TABLE_ALIGNMENT.CENTER
            table_rtl(table)
            fig_cell, sol_cell = table.rows[0].cells
            fig_cell.width = Cm(7.6)
            sol_cell.width = Cm(10.0)
            cell_shade(fig_cell, "FBFCFF")
            cell_shade(sol_cell, "FFFFFF")

            path, W, H = figs[q["id"]]
            w_cm = W / 2.0 / 96 * 2.54
            h_cm = H / 2.0 / 96 * 2.54
            s = min(7.0 / w_cm, 7.0 / h_cm, 1.5)
            fp = fig_cell.paragraphs[0]
            rtl(fp, "center")
            fp.add_run().add_picture(path, width=Cm(w_cm * s))

            sp = sol_cell.paragraphs[0]
            rtl(sp)
            r = sp.add_run("المطلوب: ")
            r.font.bold = True
            r.font.size = Pt(12)
            r.font.color.rgb = RGBColor.from_string("2A6FDB")
            r.font.name = "Arial"
            r2 = sp.add_run(plain(q["ask"]))
            r2.font.size = Pt(12)
            r2.font.name = "Arial"

            sh = sol_cell.add_paragraph()
            rtl(sh)
            rr = sh.add_run("خطوات الحل")
            rr.font.bold = True
            rr.font.size = Pt(12)
            rr.font.color.rgb = RGBColor.from_string("2A9D8F")
            rr.font.name = "Arial"

            for k, st in enumerate(q["steps"], 1):
                stp = sol_cell.add_paragraph()
                rtl(stp)
                stp.paragraph_format.space_after = Pt(1)
                rk = stp.add_run(f"{k}. ")
                rk.font.bold = True
                rk.font.size = Pt(11.5)
                rk.font.name = "Arial"
                rv = stp.add_run(plain(st))
                rv.font.size = Pt(11.5)
                rv.font.name = "Arial"

            ap = sol_cell.add_paragraph()
            rtl(ap)
            ra = ap.add_run("الإجابة: ")
            ra.font.bold = True
            ra.font.size = Pt(12)
            ra.font.color.rgb = RGBColor.from_string("12705F")
            ra.font.name = "Arial"
            rv = ap.add_run(plain(q["answer"]))
            rv.font.bold = True
            rv.font.size = Pt(12.5)
            rv.font.name = "Arial"
            shade(ap._p.get_or_add_pPr(), "EAFAF3")

            if q.get("tip"):
                tp = sol_cell.add_paragraph()
                rtl(tp)
                rt = tp.add_run("💡 " + plain(q["tip"]))
                rt.font.size = Pt(11)
                rt.font.name = "Arial"
                rt.font.color.rgb = RGBColor.from_string("B26A12")
                shade(tp._p.get_or_add_pPr(), "FFF6FB")

            para("", size=6, space_after=2)

        if ci < len(CHAPTERS):
            doc.add_page_break()

    # ---------------- appendix
    doc.add_page_break()
    hp = para("🎬 مصادر المحتوى (فيديوهات القناة)", size=19, bold=True, color="FFFFFF")
    shade(hp._p.get_or_add_pPr(), "2A6FDB")
    para("القناة: youtube.com/@emanelsayed.online — قائمة تشغيل IQ: "
         "youtube.com/playlist?list=PLO9L9sHOfqhPGsDyK_JAvNnkP0Ua8o_6Q", size=11, color="#5b6b8a")
    for i, (t, vid, tag) in enumerate(SOURCES, 1):
        rich([(f"{i}. ", True, "#2a6fdb"), (t, False, "#1f2a44"),
              (f"  ({tag})  ", False, "#8d99ae"),
              (f"https://youtu.be/{vid}", False, "#2a6fdb")], size=11)

    hp = para("🗓️ خطة مذاكرة 7 أيام", size=19, bold=True, color="FFFFFF")
    shade(hp._p.get_or_add_pPr(), "7B4BD8")
    for d, t, x in PLAN:
        rich([(f"{d}: ", True, "#7b4bd8"), (t, True, "#1f2a44"), (f" — {x}", False, "#5b6b8a")], size=11.5)

    hp = para("نصائح يوم الامتحان", size=19, bold=True, color="FFFFFF")
    shade(hp._p.get_or_add_pPr(), "2A9D8F")
    for t in TIPS:
        para("• " + t, size=11.5, space_after=2)

    doc.save(out_path)
    return out_path


# ============================================================== EXCEL
def build_xlsx(figs, out_path):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.drawing.image import Image as XLImage
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    thin = Side(style="thin", color="D7DEEC")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    # ---------------- sheet 1: questions
    ws = wb.active
    ws.title = "أسئلة IQ"
    ws.sheet_view.rightToLeft = True

    headers = ["الباب", "السؤال", "العنوان", "الشكل (صورة)", "المطلوب", "خطوات الحل", "الإجابة", "خدعة / تنبيه"]
    widths = [22, 7, 30, 46, 34, 62, 26, 34]
    head_fill = PatternFill("solid", fgColor="2A6FDB")
    for i, (h, w) in enumerate(zip(headers, widths), 1):
        c = ws.cell(row=1, column=i, value=h)
        c.font = Font(name="Arial", size=12, bold=True, color="FFFFFF")
        c.fill = head_fill
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = border
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.row_dimensions[1].height = 26
    ws.freeze_panes = "A2"

    r = 2
    for ch in CHAPTERS:
        fill = PatternFill("solid", fgColor=ch["color"].lstrip("#"))
        light = PatternFill("solid", fgColor="F7F9FF")
        for qi, q in enumerate(ch["qs"], 1):
            ws.cell(row=r, column=1, value=f'{ch["icon"]} {ch["title"]}')
            ws.cell(row=r, column=2, value=qi)
            ws.cell(row=r, column=3, value=q["title"])
            ws.cell(row=r, column=5, value=plain(q["ask"]))
            steps_txt = "\n".join(f"{k}. {plain(s)}" for k, s in enumerate(q["steps"], 1))
            ws.cell(row=r, column=6, value=steps_txt)
            ws.cell(row=r, column=7, value="✔ " + plain(q["answer"]))
            ws.cell(row=r, column=8, value=plain(q.get("tip", "")))

            for col in (1, 3, 5, 6, 7, 8):
                cell = ws.cell(row=r, column=col)
                cell.font = Font(name="Arial", size=11,
                                 bold=(col in (1, 3, 7)),
                                 color=ch["color"].lstrip("#") if col == 1 else
                                 ("12705F" if col == 7 else "1F2A44"))
                cell.alignment = Alignment(horizontal="right", vertical="top", wrap_text=True)
                cell.fill = light
                cell.border = border
            c2 = ws.cell(row=r, column=2)
            c2.font = Font(name="Arial", size=11, bold=True)
            c2.alignment = Alignment(horizontal="center", vertical="center")
            c2.border = border
            c2.fill = light
            c7 = ws.cell(row=r, column=7)
            c7.fill = PatternFill("solid", fgColor="EAFAF3")

            # image
            path, W, H = figs[q["id"]]
            max_w, max_h = 320, 190
            s = min(max_w / W, max_h / H, 1.0)
            iw, ih = max(24, int(W * s)), max(18, int(H * s))
            img = XLImage(path)
            img.width, img.height = iw, ih
            anchor = f"D{r}"
            ws.add_image(img, anchor)

            def est_lines(text, chars):
                if not text:
                    return 1
                total = 0
                for seg in str(text).split("\n"):
                    total += max(1, math.ceil(len(seg) / chars))
                return total

            need_lines = max(est_lines(q["title"], 30),
                             est_lines(plain(q["ask"]), 34),
                             est_lines(steps_txt, 56),
                             est_lines(plain(q["answer"]), 24),
                             est_lines(plain(q.get("tip", "")), 32))
            ws.row_dimensions[r].height = max(ih * 0.78 + 8, need_lines * 14.5 + 6)
            r += 1

    ws.auto_filter.ref = f"A1:H{r - 1}"

    # ---------------- sheet 2: sources
    ws2 = wb.create_sheet("المصادر")
    ws2.sheet_view.rightToLeft = True
    heads2 = ["#", "عنوان الفيديو", "المحور", "الرابط"]
    for i, h in enumerate(heads2, 1):
        c = ws2.cell(row=1, column=i, value=h)
        c.font = Font(name="Arial", size=12, bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="7B4BD8")
        c.alignment = Alignment(horizontal="center", vertical="center")
        c.border = border
    for i, w in enumerate([6, 62, 22, 46], 1):
        ws2.column_dimensions[get_column_letter(i)].width = w
    for i, (t, vid, tag) in enumerate(SOURCES, 1):
        ws2.cell(row=i + 1, column=1, value=i).alignment = Alignment(horizontal="center")
        ws2.cell(row=i + 1, column=2, value=t)
        ws2.cell(row=i + 1, column=3, value=tag)
        cell = ws2.cell(row=i + 1, column=4, value=f"https://youtu.be/{vid}")
        cell.hyperlink = f"https://youtu.be/{vid}"
        cell.font = Font(name="Arial", size=11, color="0563C1", underline="single")
        for col in range(1, 5):
            cc = ws2.cell(row=i + 1, column=col)
            if col != 4:
                cc.font = Font(name="Arial", size=11)
            cc.alignment = Alignment(horizontal="right" if col != 1 else "center", vertical="center",
                                     wrap_text=(col == 2))
            cc.border = border
    ws2.freeze_panes = "A2"

    # ---------------- sheet 3: plan
    ws3 = wb.create_sheet("خطة المذاكرة")
    ws3.sheet_view.rightToLeft = True
    heads3 = ["اليوم", "الموضوع", "المطلوب"]
    for i, h in enumerate(heads3, 1):
        c = ws3.cell(row=1, column=i, value=h)
        c.font = Font(name="Arial", size=12, bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="2A9D8F")
        c.alignment = Alignment(horizontal="center", vertical="center")
        c.border = border
    for i, w in enumerate([12, 46, 70], 1):
        ws3.column_dimensions[get_column_letter(i)].width = w
    for i, (d, t, x) in enumerate(PLAN, 1):
        ws3.cell(row=i + 1, column=1, value=d)
        ws3.cell(row=i + 1, column=2, value=t)
        ws3.cell(row=i + 1, column=3, value=x)
        for col in range(1, 4):
            cc = ws3.cell(row=i + 1, column=col)
            cc.font = Font(name="Arial", size=11.5, bold=(col == 1),
                           color="7B4BD8" if col == 1 else "1F2A44")
            cc.alignment = Alignment(horizontal="right", vertical="center", wrap_text=True)
            cc.border = border
    start = len(PLAN) + 3
    c = ws3.cell(row=start, column=1, value="نصائح يوم الامتحان")
    c.font = Font(name="Arial", size=13, bold=True, color="FFFFFF")
    c.fill = PatternFill("solid", fgColor="E63946")
    ws3.merge_cells(start_row=start, start_column=1, end_row=start, end_column=3)
    for i, t in enumerate(TIPS, 1):
        cc = ws3.cell(row=start + i, column=1, value=f"• {t}")
        cc.font = Font(name="Arial", size=11.5)
        cc.alignment = Alignment(horizontal="right", vertical="center", wrap_text=True)
        ws3.merge_cells(start_row=start + i, start_column=1, end_row=start + i, end_column=3)

    wb.save(out_path)
    return out_path


def main():
    figs = figures()
    d = build_docx(figs, os.path.join(ROOT, "IQ-guide.docx"))
    x = build_xlsx(figs, os.path.join(ROOT, "IQ-guide.xlsx"))
    print("docx:", d, os.path.getsize(d))
    print("xlsx:", x, os.path.getsize(x))


if __name__ == "__main__":
    main()
