# -*- coding: utf-8 -*-
"""
يولّد نسخة Word (.docx) من الأسئلة الأساسية المشروحة (١٠٩ أسئلة) مع رسم كل شكل كصورة.

لماذا الأسئلة الأساسية فقط؟ لأن بنوك التدريب المولّدة (أكثر من ٧٠٠ سؤال) تجعل ملف
Word ضخمًا جدًا وبطيء الفتح؛ وهي متاحة كاملة داخل iq-guide.html.

التشغيل:  python3 src/build_docx.py
المتطلبات: python-docx و pymupdf
"""
import io
import os
import re
import sys

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
sys.path.insert(0, HERE)

import pymupdf                                     # noqa: E402
from docx import Document                          # noqa: E402
from docx.enum.text import WD_ALIGN_PARAGRAPH      # noqa: E402
from docx.oxml import OxmlElement                  # noqa: E402
from docx.oxml.ns import qn                        # noqa: E402
from docx.shared import Inches, Pt, RGBColor       # noqa: E402

from model import AR_LABELS, balance_answers       # noqa: E402
from content_a import CHAPTERS_A                   # noqa: E402
from content_b import CHAPTERS_B                   # noqa: E402
from content_c import CHAPTERS_C                   # noqa: E402
from extras import EXAM, TIPS, QUICKTABLE          # noqa: E402
from genlib import ar                              # noqa: E402

CHAPTERS = balance_answers(CHAPTERS_A + CHAPTERS_B + CHAPTERS_C)

OUT = os.path.join(ROOT, "iq-guide.docx")
FONT = "Arial"
INK = RGBColor(0x0F, 0x17, 0x2A)
BLUE = RGBColor(0x1D, 0x4E, 0xD8)
GREEN = RGBColor(0x15, 0x80, 0x3D)
AMBER = RGBColor(0xB4, 0x53, 0x09)
GREY = RGBColor(0x64, 0x74, 0x8B)

_SVG_RE = re.compile(r"<svg[^>]*>(.*?)</svg>", re.S)
_TAG_RE = re.compile(r"<[^>]+>")
_cache = {}


# ----------------------------------------------------------------- أدوات RTL
def rtl(par, align="right"):
    pPr = par._p.get_or_add_pPr()
    bidi = OxmlElement("w:bidi")
    bidi.set(qn("w:val"), "1")
    pPr.append(bidi)
    par.alignment = {"right": WD_ALIGN_PARAGRAPH.RIGHT,
                     "center": WD_ALIGN_PARAGRAPH.CENTER}[align]
    for run in par.runs:
        rPr = run._element.get_or_add_rPr()
        for tag in ("w:rtl", "w:cs"):
            el = OxmlElement(tag)
            el.set(qn("w:val"), "1")
            rPr.append(el)
    return par


def para(doc, text="", size=12, bold=False, color=INK, align="right", space_after=4):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = FONT
    run._element.rPr.rFonts.set(qn("w:cs"), FONT)
    p.paragraph_format.space_after = Pt(space_after)
    return rtl(p, align)


def shade(par, hexcolor):
    pPr = par._p.get_or_add_pPr()
    sh = OxmlElement("w:shd")
    sh.set(qn("w:val"), "clear")
    sh.set(qn("w:fill"), hexcolor)
    pPr.append(sh)
    return par


# ------------------------------------------------------- تحويل SVG إلى صورة
def svg_png(cell, px=190):
    """يحوّل خلية HTML تحوي <svg> إلى صورة PNG (bytes)، مع تخزين مؤقت."""
    m = _SVG_RE.search(cell or "")
    if not m:
        return None
    inner = m.group(1)
    key = (inner, px)
    if key in _cache:
        return _cache[key]
    svg = ('<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 100 100" '
           f'width="{px}" height="{px}"><rect width="100" height="100" fill="#ffffff"/>'
           f"{inner}</svg>")
    try:
        doc = pymupdf.open(stream=svg.encode("utf-8"), filetype="svg")
        pix = doc[0].get_pixmap(dpi=96, alpha=False)
        data = pix.tobytes("png")
    except Exception:
        return None
    _cache[key] = data
    return data


def plain(cell):
    """نص خلية بلا وسوم HTML."""
    txt = _TAG_RE.sub(" ", cell or "")
    return re.sub(r"\s+", " ", txt).strip()


def figrow(doc, cells, width=1.25):
    """صف صور في فقرة واحدة."""
    imgs = [svg_png(c) for c in cells]
    imgs = [i for i in imgs if i]
    if not imgs:
        return None
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for data in imgs:
        run = p.add_run()
        run.add_picture(io.BytesIO(data), width=Inches(width))
        p.add_run("  ")
    p.paragraph_format.space_after = Pt(3)
    return p


def options_table(doc, q):
    """جدول الاختيارات الأربعة (صور أو نصوص) من اليمين لليسار."""
    opts = q["options"]
    n = len(opts)
    table = doc.add_table(rows=2, cols=n)
    table.style = "Table Grid"
    tblPr = table._tbl.tblPr
    bidi = OxmlElement("w:bidiVisual")
    tblPr.append(bidi)
    for i, opt in enumerate(opts):
        head = table.cell(0, i).paragraphs[0]
        run = head.add_run(AR_LABELS[i])
        run.font.bold = True
        run.font.size = Pt(12)
        run.font.color.rgb = BLUE if i != q["answer"] else GREEN
        run.font.name = FONT
        run._element.rPr.rFonts.set(qn("w:cs"), FONT)
        head.alignment = WD_ALIGN_PARAGRAPH.CENTER
        body = table.cell(1, i).paragraphs[0]
        body.alignment = WD_ALIGN_PARAGRAPH.CENTER
        data = svg_png(opt) if q["okind"] == "fig" else None
        if data:
            body.add_run().add_picture(io.BytesIO(data), width=Inches(1.05))
        else:
            r = body.add_run(plain(opt))
            r.font.size = Pt(12)
            r.font.name = FONT
            r._element.rPr.rFonts.set(qn("w:cs"), FONT)
            rtl(body, "center")
    return table


# ------------------------------------------------------------------ البناء
def build():
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = FONT
    style.font.size = Pt(12)
    style.element.rPr.rFonts.set(qn("w:cs"), FONT)

    for section in doc.sections:
        section.top_margin = section.bottom_margin = Inches(0.6)
        section.left_margin = section.right_margin = Inches(0.6)
        sectPr = section._sectPr
        b = OxmlElement("w:bidi")
        b.set(qn("w:val"), "1")
        sectPr.append(b)

    # ---------------- الغلاف
    para(doc, "دليل أسئلة الذكاء (IQ)", 30, True, BLUE, "center", 6)
    para(doc, "كل أنواع الأسئلة + طريقة الحل + رسمة لكل شكل", 15, False, GREY, "center", 14)
    para(doc, "مستخلص من شروحات قناة: هنتعلم أون لاين مع إيمان السيد", 12, False, INK, "center", 2)
    para(doc, "youtube.com/@emanelsayed.online", 11, False, BLUE, "center", 16)
    total = sum(len(c["questions"]) for c in CHAPTERS)
    para(doc, f"{ar(len(CHAPTERS))} محورًا · {ar(total)} سؤالًا أساسيًا مشروحًا خطوة بخطوة", 13, True, INK, "center", 6)
    para(doc, "ملاحظة: بنوك التدريب الإضافية (أكثر من ٧٠٠ سؤال مولّد) موجودة كاملة "
              "في ملف iq-guide.html المرفق؛ هذا الملف يحتوي الأسئلة الأساسية المشروحة.",
         11, False, GREY, "center", 10)

    # ---------------- فهرس مختصر
    para(doc, "المحتويات", 18, True, BLUE, "right", 6)
    for ch in CHAPTERS:
        para(doc, f"{ar(ch['num'])}. {ch['title']} — {ar(len(ch['questions']))} أسئلة", 12, False, INK, "right", 1)
    doc.add_page_break()

    # ---------------- المحاور
    for ch in CHAPTERS:
        head = para(doc, f"{ar(ch['num'])}. {ch['icon']} {ch['title']}", 20, True, BLUE, "right", 2)
        shade(head, "EFF6FF")
        para(doc, ch["subtitle"], 11, False, GREY, "right", 8)

        para(doc, "فكرة السؤال", 13, True, AMBER, "right", 2)
        para(doc, ch["idea"], 12, False, INK, "right", 8)

        para(doc, "طريقة الحل خطوة بخطوة", 13, True, AMBER, "right", 2)
        for i, step in enumerate(ch["how"], 1):
            para(doc, f"{ar(i)}. {step}", 12, False, INK, "right", 1)

        para(doc, "أخطاء شائعة", 13, True, AMBER, "right", 2)
        for trap in ch["traps"]:
            para(doc, f"• {trap}", 12, False, INK, "right", 1)

        if ch.get("source"):
            para(doc, "من فيديوهات القناة: " + " | ".join(ch["source"]), 10, False, GREY, "right", 8)

        for qi, q in enumerate(ch["questions"], 1):
            para(doc, f"سؤال {ar(ch['num'])}-{ar(qi)}: {q['prompt']}", 13, True, INK, "right", 3)
            if q.get("stem"):
                figrow(doc, q["stem"])
                if q.get("stem_note"):
                    para(doc, q["stem_note"], 10, False, GREY, "center", 4)
            options_table(doc, q)
            ans = AR_LABELS[q["answer"]]
            atxt = q.get("answer_text") or (plain(q["options"][q["answer"]])
                                            if q["okind"] == "text" else "")
            line = f"✔ الإجابة الصحيحة: ({ans})" + (f" — {atxt}" if atxt else "")
            para(doc, line, 12, True, GREEN, "right", 2)
            for i, step in enumerate(q["steps"], 1):
                para(doc, f"{ar(i)}. {step}", 11, False, INK, "right", 1)
            if q.get("tip"):
                para(doc, f"⚠ انتبه: {q['tip']}", 11, False, AMBER, "right", 6)
            else:
                para(doc, "", 6, False, INK, "right", 6)
        doc.add_page_break()

    # ---------------- جدول الأنماط السريع
    para(doc, "جدول الأنماط السريع", 20, True, BLUE, "right", 6)
    t = doc.add_table(rows=1, cols=3)
    t.style = "Table Grid"
    t._tbl.tblPr.append(OxmlElement("w:bidiVisual"))
    hdr = ["لو رأيت هذا في الامتحان…", "فهذا النمط", "افعل هذا فورًا"]
    for i, h in enumerate(hdr):
        p = t.cell(0, i).paragraphs[0]
        r = p.add_run(h)
        r.font.bold = True
        r.font.name = FONT
        r._element.rPr.rFonts.set(qn("w:cs"), FONT)
        rtl(p, "center")
    for row in QUICKTABLE:
        cells = t.add_row().cells
        for i, val in enumerate(row[:3]):
            p = cells[i].paragraphs[0]
            r = p.add_run(plain(str(val)))
            r.font.size = Pt(10)
            r.font.name = FONT
            r._element.rPr.rFonts.set(qn("w:cs"), FONT)
            rtl(p, "right")
    doc.add_page_break()

    # ---------------- نصائح
    para(doc, "١٠ نصائح قبل الامتحان", 20, True, BLUE, "right", 6)
    for i, (title, body) in enumerate(TIPS, 1):
        para(doc, f"{ar(i)}. {plain(title)}", 12, True, INK, "right", 1)
        para(doc, plain(body), 11, False, GREY, "right", 4)
    doc.add_page_break()

    # ---------------- الاختبار الذاتي
    para(doc, "اختبار ذاتي (١٠ أسئلة)", 20, True, BLUE, "right", 6)
    answers = []
    for i, (prompt, opts, ansi, why) in enumerate(EXAM, 1):
        para(doc, f"{ar(i)}. {plain(prompt)}", 12, True, INK, "right", 2)
        letters = "   ".join(f"({AR_LABELS[j]}) {plain(o)}" for j, o in enumerate(opts))
        para(doc, letters, 12, False, INK, "right", 2)
        answers.append(f"{ar(i)}) ({AR_LABELS[ansi]}) {plain(opts[ansi])}")
    para(doc, "الإجابات مع التعليل", 14, True, GREEN, "right", 4)
    for i, (prompt, opts, ansi, why) in enumerate(EXAM, 1):
        para(doc, f"{ar(i)}) ({AR_LABELS[ansi]}) {plain(opts[ansi])} — {plain(why)}",
             11, False, INK, "right", 2)

    doc.save(OUT)
    size = os.path.getsize(OUT) // 1024
    print(f"✅ تم إنشاء {OUT}  ({size} KB) — {len(CHAPTERS)} محورًا و{total} سؤالًا أساسيًا")


if __name__ == "__main__":
    build()
